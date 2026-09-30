import tempfile
from pathlib import Path

from openpyxl import Workbook

import building_lookup_app as b

tmp = Path(tempfile.mkdtemp())
xlsx = tmp / "abc123_book.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "First"
ws.append(["lat", "lon", "id"])
ws.append(["1.1", "2.2", "a"])
ws.append(["3.3", "4.4", "b"])
ws2 = wb.create_sheet("Sec'ond")
ws2.append(["lat", "lon", "extra"])
ws2.append(["5.5", "6.6", "x"])
wb.create_sheet("Empty")
wb.save(xlsx)

sheets = b.list_excel_sheet_names(xlsx)
print("sheets", sheets)
print("preview first", b.preview_excel_file(xlsx))
print("find before", b.find_upload(tmp, "abc123").name)

csv = b.excel_sheet_csv_path(xlsx)
b.convert_excel_sheets_to_csv(xlsx, csv, [sheets[1]])
print("find after", b.find_upload(tmp, "abc123").name)
print("sheet2", b.preview_uploaded_file(csv))

b.convert_excel_sheets_to_csv(xlsx, csv, sheets)
print("ALL", b.preview_uploaded_file(csv))
print("df", b.load_csv_dataframe(csv).to_dict("records"))
